"""
Document Loader Service for Phase 3.
Reuses and extends Phase 2 document loading and page-tracking logic.
Extracts structured text, page metadata, and statistics from PDF, DOCX, TXT, CSV, and JSON files.
Prepared for integration into the Phase 3 RAG chunking and vector storage pipeline.
"""

import os
import csv
import io
import json
from typing import Dict, Any, List, Optional


def load_pdf(file_path: str) -> Dict[str, Any]:
    """
    Extract text and page-level metadata from a PDF file using pypdf.
    Reused and hardened from Phase 2 document_loader.
    """
    filename = os.path.basename(file_path)
    result: Dict[str, Any] = {
        "filename": filename,
        "filepath": file_path,
        "file_type": "PDF",
        "text": "",
        "pages": [],
        "num_pages": 0,
        "char_count": 0,
        "word_count": 0,
        "status": "success",
        "error_message": None,
    }

    try:
        import pypdf

        reader = pypdf.PdfReader(file_path)

        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                result["status"] = "error"
                result["error_message"] = "PDF is password-protected or encrypted."
                return result

        pages_data: List[Dict[str, Any]] = []
        full_text_parts: List[str] = []

        for idx, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            page_text_clean = page_text.strip()
            pages_data.append({
                "page_number": idx,
                "text": page_text_clean,
                "char_count": len(page_text_clean),
                "word_count": len(page_text_clean.split()),
            })
            if page_text_clean:
                full_text_parts.append(page_text_clean)

        combined_text = "\n\n".join(full_text_parts)
        result["text"] = combined_text
        result["pages"] = pages_data
        result["num_pages"] = max(1, len(reader.pages))
        result["char_count"] = len(combined_text)
        result["word_count"] = len(combined_text.split())

        if not combined_text.strip():
            result["status"] = "warning"
            result["error_message"] = "PDF contains no extractable text (it may be a scanned image)."

    except Exception as e:
        result["status"] = "error"
        result["error_message"] = f"Failed to extract PDF text: {str(e)}"

    return result


def load_docx(file_path: str) -> Dict[str, Any]:
    """
    Extract text from a DOCX document using python-docx.
    Reused from Phase 2 document_loader.
    """
    filename = os.path.basename(file_path)
    result: Dict[str, Any] = {
        "filename": filename,
        "filepath": file_path,
        "file_type": "DOCX",
        "text": "",
        "pages": [],
        "num_pages": 1,
        "char_count": 0,
        "word_count": 0,
        "status": "success",
        "error_message": None,
    }

    try:
        import docx

        doc = docx.Document(file_path)

        paragraphs_text: List[str] = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                paragraphs_text.append(text)

        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    paragraphs_text.append(" | ".join(row_cells))

        combined_text = "\n\n".join(paragraphs_text)
        result["text"] = combined_text
        result["pages"] = [{
            "page_number": 1,
            "text": combined_text,
            "char_count": len(combined_text),
            "word_count": len(combined_text.split()),
        }]
        result["num_pages"] = 1
        result["char_count"] = len(combined_text)
        result["word_count"] = len(combined_text.split())

        if not combined_text.strip():
            result["status"] = "warning"
            result["error_message"] = "DOCX document is empty."

    except Exception as e:
        result["status"] = "error"
        result["error_message"] = f"Failed to extract DOCX text: {str(e)}"

    return result


def load_txt(file_path: str) -> Dict[str, Any]:
    """
    Extract text from a plain text file.
    Reused from Phase 2 document_loader.
    """
    filename = os.path.basename(file_path)
    result: Dict[str, Any] = {
        "filename": filename,
        "filepath": file_path,
        "file_type": "TXT",
        "text": "",
        "pages": [],
        "num_pages": 1,
        "char_count": 0,
        "word_count": 0,
        "status": "success",
        "error_message": None,
    }

    try:
        content = ""
        for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                with open(file_path, "r", encoding=encoding) as f:
                    content = f.read()
                break
            except (UnicodeDecodeError, UnicodeError):
                continue

        combined_text = content.strip()
        result["text"] = combined_text
        result["pages"] = [{
            "page_number": 1,
            "text": combined_text,
            "char_count": len(combined_text),
            "word_count": len(combined_text.split()),
        }]
        result["num_pages"] = 1
        result["char_count"] = len(combined_text)
        result["word_count"] = len(combined_text.split())

        if not combined_text:
            result["status"] = "warning"
            result["error_message"] = "TXT file is empty."

    except Exception as e:
        result["status"] = "error"
        result["error_message"] = f"Failed to read TXT file: {str(e)}"

    return result


def load_csv(file_path: str) -> Dict[str, Any]:
    """
    Extract structured tabular text from a CSV file.
    Tracks row counts and formats data cleanly for downstream RAG chunking.
    """
    filename = os.path.basename(file_path)
    result: Dict[str, Any] = {
        "filename": filename,
        "filepath": file_path,
        "file_type": "CSV",
        "text": "",
        "pages": [],
        "num_pages": 1,
        "char_count": 0,
        "word_count": 0,
        "status": "success",
        "error_message": None,
    }

    try:
        raw_text = ""
        for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                with open(file_path, "r", encoding=encoding) as f:
                    raw_text = f.read()
                break
            except (UnicodeDecodeError, UnicodeError):
                continue

        if not raw_text.strip():
            result["status"] = "warning"
            result["error_message"] = "CSV file is empty."
            return result

        # Detect dialect / delimiter
        try:
            sniffer = csv.Sniffer()
            dialect = sniffer.sniff(raw_text[:2048])
            delimiter = dialect.delimiter
        except Exception:
            delimiter = ","

        reader = csv.reader(io.StringIO(raw_text), delimiter=delimiter)
        rows = [row for row in reader if any(cell.strip() for cell in row)]

        if not rows:
            result["status"] = "warning"
            result["error_message"] = "CSV file contains no valid rows."
            return result

        header = rows[0]
        formatted_lines: List[str] = [f"Columns: {', '.join(header)}"]

        if len(rows) > 1:
            for r_idx, row in enumerate(rows[1:], start=1):
                cell_pairs = []
                for c_idx, cell in enumerate(row):
                    col_name = header[c_idx] if c_idx < len(header) else f"Col{c_idx + 1}"
                    cell_pairs.append(f"{col_name}: {cell.strip()}")
                formatted_lines.append(f"Row {r_idx}: " + " | ".join(cell_pairs))

        combined_text = "\n".join(formatted_lines)
        # Allocate ~50 rows per virtual page for RAG source page referencing
        num_virtual_pages = max(1, (len(rows) + 49) // 50)

        # Break into page chunks
        pages_data = []
        rows_per_page = 50
        for p in range(num_virtual_pages):
            start = p * rows_per_page
            end = min(len(formatted_lines), (p + 1) * rows_per_page)
            p_text = "\n".join(formatted_lines[start:end])
            pages_data.append({
                "page_number": p + 1,
                "text": p_text,
                "char_count": len(p_text),
                "word_count": len(p_text.split()),
            })

        result["text"] = combined_text
        result["pages"] = pages_data
        result["num_pages"] = num_virtual_pages
        result["char_count"] = len(combined_text)
        result["word_count"] = len(combined_text.split())

    except Exception as e:
        result["status"] = "error"
        result["error_message"] = f"Failed to parse CSV file: {str(e)}"

    return result


def load_json(file_path: str) -> Dict[str, Any]:
    """
    Extract structured text from a JSON file.
    Validates JSON integrity and formats it legibly for RAG embedding.
    """
    filename = os.path.basename(file_path)
    result: Dict[str, Any] = {
        "filename": filename,
        "filepath": file_path,
        "file_type": "JSON",
        "text": "",
        "pages": [],
        "num_pages": 1,
        "char_count": 0,
        "word_count": 0,
        "status": "success",
        "error_message": None,
    }

    try:
        raw_text = ""
        for encoding in ["utf-8", "utf-8-sig", "latin-1"]:
            try:
                with open(file_path, "r", encoding=encoding) as f:
                    raw_text = f.read()
                break
            except (UnicodeDecodeError, UnicodeError):
                continue

        if not raw_text.strip():
            result["status"] = "warning"
            result["error_message"] = "JSON file is empty."
            return result

        parsed = json.loads(raw_text)
        formatted_text = json.dumps(parsed, indent=2, ensure_ascii=False)

        result["text"] = formatted_text
        result["pages"] = [{
            "page_number": 1,
            "text": formatted_text,
            "char_count": len(formatted_text),
            "word_count": len(formatted_text.split()),
        }]
        result["num_pages"] = 1
        result["char_count"] = len(formatted_text)
        result["word_count"] = len(formatted_text.split())

    except json.JSONDecodeError as e:
        result["status"] = "error"
        result["error_message"] = f"Invalid JSON syntax: {str(e)}"
    except Exception as e:
        result["status"] = "error"
        result["error_message"] = f"Failed to parse JSON file: {str(e)}"

    return result


def load_document(file_path: str) -> Dict[str, Any]:
    """
    Universal document loader: dispatches to appropriate parser based on file extension.
    Supports PDF, DOCX, TXT, CSV, and JSON.
    """
    if not os.path.exists(file_path):
        return {
            "filename": os.path.basename(file_path),
            "filepath": file_path,
            "file_type": "UNKNOWN",
            "text": "",
            "pages": [],
            "num_pages": 0,
            "char_count": 0,
            "word_count": 0,
            "status": "error",
            "error_message": f"File does not exist: {file_path}",
        }

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return load_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        return load_docx(file_path)
    elif ext == ".csv":
        return load_csv(file_path)
    elif ext == ".json":
        return load_json(file_path)
    elif ext in [".txt", ".md"]:
        return load_txt(file_path)
    else:
        return {
            "filename": os.path.basename(file_path),
            "filepath": file_path,
            "file_type": ext.upper().lstrip(".") or "UNKNOWN",
            "text": "",
            "pages": [],
            "num_pages": 0,
            "char_count": 0,
            "word_count": 0,
            "status": "error",
            "error_message": f"Unsupported file extension '{ext}'. Supported: PDF, DOCX, TXT, CSV, JSON.",
        }
